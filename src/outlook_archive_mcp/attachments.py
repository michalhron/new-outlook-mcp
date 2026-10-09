"""Locate attachment bytes on disk and extract text from them. Never downloads anything."""

from __future__ import annotations

import csv
import io
import mimetypes
import re
import zlib
from dataclasses import dataclass
from pathlib import Path

from . import mime, paths
from .db import Archive

NOT_CACHED = (
    "This attachment is not stored on this Mac. Open the message in Outlook so it downloads "
    "the attachment, then run a sync (sync_now) and try again."
)
#: Inline images below this size are treated as signature logos and hidden by default.
SMALL_INLINE_IMAGE = 100_000
TEXT_TYPES = {"text/plain", "text/csv", "text/calendar", "text/markdown", "application/json", "text/x-vcard",
              "text/vcard", "application/ics", "text/tab-separated-values", "application/xml", "text/xml"}
TEXT_EXTS = {".txt", ".csv", ".tsv", ".ics", ".vcf", ".md", ".json", ".xml", ".log", ".eml"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".heic", ".tif", ".tiff", ".bmp", ".webp"}


class AttachmentError(ValueError):
    pass


@dataclass
class AttachmentRow:
    id: int
    message_pk: int
    source: str
    filename: str | None
    content_type: str | None
    size: int | None
    content_id: str | None
    is_inline: bool
    local_path: str | None
    storage: str | None
    part_index: int | None

    @classmethod
    def from_row(cls, r) -> "AttachmentRow":
        return cls(r["id"], r["message_pk"], r["source"], r["filename"], r["content_type"], r["size"],
                   r["content_id"], bool(r["is_inline"]), r["local_path"], r["storage"], r["part_index"])

    @property
    def display_name(self) -> str:
        if self.filename:
            return self.filename
        ext = mimetypes.guess_extension(self.content_type or "") or ".bin"
        return f"attachment-{self.id}{ext}"

    @property
    def is_small_inline_image(self) -> bool:
        return (self.is_inline and (self.content_type or "").startswith("image/")
                and (self.size or 0) < SMALL_INLINE_IMAGE)


def _has_raw(archive: Archive, message_pk: int) -> bool:
    row = archive.conn.execute("SELECT raw_source_z IS NOT NULL FROM messages WHERE id = ?",
                               (message_pk,)).fetchone()
    return bool(row and row[0])


def _file_ok(att: AttachmentRow) -> bool:
    return bool(att.local_path) and Path(att.local_path).is_file()


def available(archive: Archive, att: AttachmentRow) -> bool:
    if att.storage == "file":
        return _file_ok(att)
    if att.storage == "mime_file":
        # A part of a full message source can also be cut from the copy stored in the archive.
        return _file_ok(att) or (att.part_index is not None and _has_raw(archive, att.message_pk))
    if att.storage == "raw_mime":
        return _has_raw(archive, att.message_pk)
    return False


def get_row(archive: Archive, attachment_id: int | str) -> AttachmentRow:
    try:
        aid = int(attachment_id)
    except (TypeError, ValueError) as exc:
        raise AttachmentError(f"invalid attachment id {attachment_id!r}") from exc
    r = archive.conn.execute("SELECT * FROM attachments WHERE id = ?", (aid,)).fetchone()
    if r is None:
        raise AttachmentError(f"no attachment with id {attachment_id!r}")
    return AttachmentRow.from_row(r)


def _mime_part_bytes(data: bytes, part_index: int | None) -> bytes:
    msg = mime.parse_message_bytes(data)
    if part_index is None:
        payload = msg.get_payload(decode=True)
        if payload is None:
            raise AttachmentError("attachment part has no decodable content")
        return payload
    for i, part in enumerate(mime.iter_attachment_parts(msg)):
        if i == part_index:
            payload = part.get_payload(decode=True)
            if payload is None:
                raise AttachmentError("attachment part has no decodable content")
            return payload
    raise AttachmentError("attachment part not found in message source")


def read_bytes(archive: Archive, att: AttachmentRow) -> bytes:
    if not available(archive, att):
        raise AttachmentError(NOT_CACHED)
    if att.storage == "file":
        return Path(att.local_path).read_bytes()
    if att.storage == "mime_file" and _file_ok(att):
        return _mime_part_bytes(Path(att.local_path).read_bytes(), att.part_index)
    raw = archive.conn.execute("SELECT raw_source_z FROM messages WHERE id = ?", (att.message_pk,)).fetchone()[0]
    return _mime_part_bytes(zlib.decompress(raw), att.part_index)


def _safe_name(name: str) -> str:
    name = re.sub(r"[/\\\x00-\x1f:]", "_", name).strip(". ") or "attachment"
    return name[:200]


def materialize(archive: Archive, att: AttachmentRow) -> Path:
    """Return a real file path for the attachment.

    Plain cached files are returned as they are. Attachments stored inside MIME
    are decoded once into our own cache directory (never into Outlook's folders).
    """
    if att.storage == "file":
        if not available(archive, att):
            raise AttachmentError(NOT_CACHED)
        return Path(att.local_path)
    target = paths.app_dir() / "attachments" / str(att.id) / _safe_name(att.display_name)
    if not target.exists():
        data = read_bytes(archive, att)
        target.parent.mkdir(parents=True, exist_ok=True)
        tmp = target.with_suffix(target.suffix + ".part")
        tmp.write_bytes(data)
        tmp.replace(target)
    return target


# ------------------------------------------------------------- text extraction

def _kind(att: AttachmentRow) -> str:
    ext = Path(att.display_name).suffix.lower()
    ctype = (att.content_type or "").lower()
    if ext == ".pdf" or ctype == "application/pdf":
        return "pdf"
    if ext == ".docx" or ctype.endswith("wordprocessingml.document"):
        return "docx"
    if ext in (".xlsx", ".xlsm") or ctype.endswith("spreadsheetml.sheet"):
        return "xlsx"
    if ext in (".html", ".htm") or ctype == "text/html":
        return "html"
    if ext in TEXT_EXTS or ctype in TEXT_TYPES or ctype.startswith("text/"):
        return "text"
    if ext in IMAGE_EXTS or ctype.startswith("image/"):
        return "image"
    return "other"


def _decode_text(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append(f"--- page {i} ---\n{page.extract_text() or ''}")
    return "\n".join(pages)


def _docx_text(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append("\t".join(c.text for c in row.cells))
    return "\n".join(out)


def _xlsx_text(path: Path, max_rows: int = 5000) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            out.append(f"--- sheet {ws.title} ---")
            buf = io.StringIO()
            w = csv.writer(buf)
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= max_rows:
                    buf.write(f"... truncated after {max_rows} rows\n")
                    break
                w.writerow(["" if v is None else v for v in row])
            out.append(buf.getvalue().rstrip())
    finally:
        wb.close()
    return "\n".join(out)


def extract_text(archive: Archive, att: AttachmentRow) -> tuple[str | None, str]:
    """Return (text, kind). Text is None for images and unsupported types."""
    kind = _kind(att)
    if kind in ("image", "other"):
        return None, kind
    if kind == "text":
        return _decode_text(read_bytes(archive, att)), kind
    if kind == "html":
        return mime.html_to_text(_decode_text(read_bytes(archive, att))), kind
    path = materialize(archive, att)
    try:
        if kind == "pdf":
            return _pdf_text(path), kind
        if kind == "docx":
            return _docx_text(path), kind
        return _xlsx_text(path), kind
    except ImportError as exc:
        raise AttachmentError(f"text extraction for {kind} needs an extra package: {exc.name}") from exc
    except Exception as exc:
        raise AttachmentError(f"could not extract text from {att.display_name}: {exc}") from exc
