from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path

import pytest

from outlook_archive_mcp import tools
from outlook_archive_mcp.model import AttachmentInfo, MessageRecord


def _minimal_pdf(text: str) -> bytes:
    """A one-page PDF with a text line, written by hand (no PDF library needed)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


@pytest.fixture
def files_archive(archive, tmp_path):
    """One message whose attachments are plain files, as in New Outlook's Files/ cache."""
    import docx
    import openpyxl

    files = tmp_path / "Files"
    files.mkdir()
    (files / "minutes.pdf").write_bytes(_minimal_pdf("Minutes of the walrus committee"))
    d = docx.Document()
    d.add_paragraph("Draft chapter about lighthouses")
    d.save(files / "chapter.docx")
    wb = openpyxl.Workbook()
    wb.active.title = "Costs"
    wb.active.append(["item", "eur"])
    wb.active.append(["ferry", 120])
    wb.save(files / "costs.xlsx")
    (files / "invite.ics").write_text("BEGIN:VCALENDAR\nSUMMARY:Seminar\nEND:VCALENDAR\n")
    (files / "logo.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 100)
    (files / "photo.jpg").write_bytes(b"\xff\xd8\xff" + b"\x00" * 100)

    def f(name, ctype, inline=False, size=None, path=True):
        p = files / name
        return AttachmentInfo(filename=name, content_type=ctype, size=size or (p.stat().st_size if p.exists() else 1),
                              is_inline=inline, local_path=str(p) if path else None,
                              storage="file" if path else None)

    rec = MessageRecord(
        source="hxstore", source_key="m1", message_id="<files-1@example.org>", subject="Committee papers",
        date=datetime(2026, 9, 1, tzinfo=timezone.utc), body_text="See attached.",
        attachments=[
            f("minutes.pdf", "application/pdf"),
            f("chapter.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            f("costs.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            f("invite.ics", "text/calendar"),
            f("logo.png", "image/png", inline=True),
            f("photo.jpg", "image/jpeg"),
            f("not-downloaded.pdf", "application/pdf", size=5000, path=False),
        ],
    )
    with archive.transaction():
        archive.upsert(rec)
    return archive


def _ids(archive, email_id="files-1@example.org", **kw):
    return {a["filename"]: a for a in tools.list_attachments(archive, email_id, **kw)["attachments"]}


def test_list_attachments_hides_small_inline_images(files_archive):
    res = tools.list_attachments(files_archive, "files-1@example.org")
    names = [a["filename"] for a in res["attachments"]]
    assert "logo.png" not in names and "photo.jpg" in names
    assert res["hidden_inline_images"] == 1
    assert "logo.png" in _ids(files_archive, include_inline=True)


def test_not_cached_attachment_is_reported(files_archive):
    a = _ids(files_archive)["not-downloaded.pdf"]
    assert a["available_locally"] is False and "Open the message in Outlook" in a["note"]
    with pytest.raises(tools.ToolInputError, match="not stored on this Mac"):
        tools.get_attachment(files_archive, a["attachment_id"], mode="text")


@pytest.mark.parametrize("name, needle", [
    ("minutes.pdf", "walrus committee"),
    ("chapter.docx", "lighthouses"),
    ("costs.xlsx", "ferry,120"),
    ("invite.ics", "SUMMARY:Seminar"),
])
def test_text_extraction(files_archive, name, needle):
    att = _ids(files_archive)[name]
    res = tools.get_attachment(files_archive, att["attachment_id"], mode="text")
    assert needle in res["text"]


def test_text_paging(files_archive):
    att = _ids(files_archive)["costs.xlsx"]
    first = tools.get_attachment(files_archive, att["attachment_id"], mode="text", max_chars=100)
    full = tools.get_attachment(files_archive, att["attachment_id"], mode="text")
    if first["truncated"]:
        rest = tools.get_attachment(files_archive, att["attachment_id"], offset=first["next_offset"])
        assert first["text"] + rest["text"] == full["text"]
    else:
        assert first["text"] == full["text"]


def test_image_returns_path(files_archive):
    att = _ids(files_archive)["photo.jpg"]
    res = tools.get_attachment(files_archive, att["attachment_id"], mode="text")
    assert res["kind"] == "image" and Path(res["path"]).name == "photo.jpg" and "text" not in res


def test_open_mode_uses_opener(files_archive):
    att = _ids(files_archive)["minutes.pdf"]
    opened = []
    res = tools.get_attachment(files_archive, att["attachment_id"], mode="open", opener=opened.append)
    assert res["opened"] and opened == [res["path"]]


def test_attachment_name_filter(files_archive, loaded):
    assert tools.search_emails(files_archive, None, attachment_name="chapter")["total"] == 1
    assert tools.search_emails(files_archive, None, attachment_name="budget.csv")["total"] == 1
    assert tools.search_emails(files_archive, None, attachment_name="nothing-like-this")["total"] == 0


def test_legacy_mime_attachments(loaded):
    # From the full RFC 822 source (MSrc block): part of a multipart message.
    att = _ids(loaded, "alpha-1@example.org")["budget.csv"]
    assert att["available_locally"] and att["source"] == "legacy"
    res = tools.get_attachment(loaded, att["attachment_id"], mode="text")
    assert "zebra,42" in res["text"]
    # From a standalone .olk15MsgAttachment block (one MIME part, base64, CR line endings).
    att = _ids(loaded, "beta-1@example.com")["report.pdf"]
    p = Path(tools.get_attachment(loaded, att["attachment_id"], mode="path")["path"])
    assert p.name == "report.pdf" and p.read_bytes().startswith(b"%PDF-1.4")
    assert "outlook" not in str(p.parent).lower()  # decoded into our own cache dir


def test_legacy_source_attachment_survives_missing_data_folder(loaded, legacy_data):
    import shutil

    att = _ids(loaded, "alpha-1@example.org")["budget.csv"]
    shutil.rmtree(legacy_data)
    res = tools.get_attachment(loaded, att["attachment_id"], mode="text")
    assert "zebra,42" in res["text"]  # cut from the raw source stored in the archive


def test_dedup_across_sources_prefers_available(files_archive):
    # The same message seen by a second source without the cached file.
    rec = MessageRecord(source="legacy", source_key="77", message_id="files-1@example.org",
                        attachments=[AttachmentInfo(filename="minutes.pdf", content_type="application/pdf")])
    with files_archive.transaction():
        files_archive.upsert(rec)
    att = [a for a in tools.list_attachments(files_archive, "files-1@example.org")["attachments"]
           if a["filename"] == "minutes.pdf"]
    assert len(att) == 1 and att[0]["available_locally"] and att[0]["source"] == "hxstore"
